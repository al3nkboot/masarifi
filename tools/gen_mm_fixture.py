# Generates a SYNTHETIC Money Manager export fixture (fake data) with the same column layout as the real
# Realbyte "Money Manager" Excel export: Date | Account | Category | Subcategory | Note | OMR | Income/Expense |
# Description | Amount | Currency | Account
import csv, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Border, Side

HEADER = ["Date", "Account", "Category", "Subcategory", "Note", "OMR", "Income/Expense", "Description", "Amount", "Currency", "Account"]
D = dt.datetime
ROWS = [
    # date, account, category, sub, note, omr, type, desc, amount, currency
    (D(2024, 3, 31, 21, 5, 10, 250000), "Cash", "Food", "", "شاورما", 1.5, "Expense", "", "1.5", "OMR"),
    (D(2024, 3, 30, 9, 0, 0), "Bank Muscat", "Salary", "", "راتب مارس", 800.0, "Income", "", "800", "OMR"),
    (D(2024, 3, 28), "Cash", "Grecory", "", "", 12.25, "Expense", "", "12.25", "OMR"),            # date-only, unknown category
    (D(2024, 3, 20, 18, 30, 0), "Bank Muscat", "Cash", "", "ATM", 50.0, "Transfer-Out", "", "50", "OMR"),
    (D(2024, 3, 20, 18, 30, 0), "Cash", "Bank Muscat", "", "ATM", 50.0, "Transfer-In", "", "50", "OMR"),
    (D(2024, 3, 15, 13, 45, 30, 500000), "Visa Card", "Electronics ", "Phone", "سماعة", 25.75, "Expense", "Online order", "25.75", "OMR"),
    (D(2024, 3, 10, 8, 15, 0), "Visa Card", "Trip", "", "Hotel", 38.45, "Expense", "", "100", "USD"),  # foreign currency
    (D(2024, 3, 5, 7, 0, 0), "Savings", "Bank Muscat", "", "", 20.0, "Transfer-Out", "", "20", "OMR"),  # unpaired transfer-out
    (D(2024, 2, 29, 23, 59, 59), "Cash", "fuel ", "", "Shell", 7.2, "Expense", "", "7.2", "OMR"),
    (D(2024, 2, 14, 12, 0, 0), "Cash", "Gift", "", "ورد", 5.0, "Expense", "", "5", "OMR"),
    (D(2024, 2, 1, 10, 0, 0), "Bank Muscat", "Bonus", "", "", 100.0, "Income", "", "100", "OMR"),
    (D(2024, 2, 1, 10, 0, 0), "Cash", "Food", "", "duplicate-looking", 2.0, "Expense", "", "2", "OMR"),
    (D(2024, 2, 1, 10, 0, 0), "Cash", "Food", "", "duplicate-looking", 2.0, "Expense", "", "2", "OMR"),  # identical row → both kept
    ("not a date", "Cash", "Food", "", "bad date", 1.0, "Expense", "", "1", "OMR"),
    (D(2024, 1, 5, 10, 0, 0), "Cash", "Food", "", "zero", 0.0, "Expense", "", "0", "OMR"),
    (D(2024, 1, 3, 10, 0, 0), "Cash", "Other", "", "weird", 3.0, "Modified Bal.", "", "3", "OMR"),
]

here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, "..", "app", "src", "test", "resources", "mm")
os.makedirs(out, exist_ok=True)
wb = Workbook(); ws = wb.active; ws.title = "Money Manager"
ws.append(HEADER)
thin = Side(style="thin")
for c in ws[1]:
    c.fill = PatternFill("solid", fgColor="C0C0C0"); c.alignment = Alignment(horizontal="center"); c.border = Border(thin, thin, thin, thin)
for r in ROWS:
    d, acc, cat, sub, note, omr, typ, desc, amt, cur = r
    ws.append([d, acc, cat, sub, note, omr, typ, desc, amt, cur, omr])
    cell = ws.cell(row=ws.max_row, column=1)
    if isinstance(d, D):
        cell.number_format = "MM/dd/yyyy" if (d.hour, d.minute, d.second) == (0, 0, 0) else "MM/dd/yyyy HH:mm:ss"
wb.save(os.path.join(out, "money_manager_sample.xlsx"))

with open(os.path.join(out, "money_manager_sample.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(HEADER)
    for r in ROWS:
        d, acc, cat, sub, note, omr, typ, desc, amt, cur = r
        ds = d.strftime("%m/%d/%Y %H:%M:%S") if isinstance(d, D) else d
        w.writerow([ds, acc, cat, sub, note, omr, typ, desc, amt, cur, omr])
print("fixture rows", len(ROWS))
